#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《直线绘制方式扩展-深度思考》docx。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Microsoft YaHei"
DOC = "D:/01-CAD/02-avticationCAD/sketch-cad-cpp/docs/直线绘制方式扩展-深度思考.docx"

os.makedirs(os.path.dirname(DOC), exist_ok=True)
doc = Document()

# 默认字体
style = doc.styles['Normal']
style.font.name = FONT
style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
style.font.size = Pt(10.5)

# 标题
T = doc.add_heading("航空CAD-草图绘制模块-深度思考", level=0)
T.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = T.runs[0]
r.font.name = FONT
r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
r.font.size = Pt(18)
r.font.bold = True

SUB = doc.add_paragraph("直线绘制方式扩展：根据菜单实现 2 Points / Horizontal / Vertical / Parallel / Tangent / Snake 等 18 种画法")
SUB.alignment = WD_ALIGN_PARAGRAPH.CENTER
SUB.runs[0].font.name = FONT
SUB.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
SUB.runs[0].font.size = Pt(12)
SUB.runs[0].font.color.rgb = RGBColor(0x40, 0x40, 0x40)

def P(text, bold=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    return p

def H(text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.color.rgb = RGBColor(0x1A, 0x1D, 0x28)
    return h

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return p

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn('w:eastAsia'), "Consolas")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x2B, 0x59, 0xA8)
    return p

# ===================== 正文 =====================
H("一、需求理解", 1)
P("用户提供的附图是一张典型的 CAD 直线绘制下拉菜单，包含 18 个可归类为「直线画法」的条目：")
bullet("基础两点：2 Points")
bullet("正交约束：Horizontal、Vertical、Orthogonal、Orthogonal From Line、From Point To Line")
bullet("角度约束：Angle、Relative angle、Angle From Line")
bullet("平行约束：Parallel、Parallel through point")
bullet("切线族：Tangent (P,C)、Tangent (C,C)、Tangent Orthogonal")
bullet("构造线：Bisector")
bullet("正交折线：Snake、Snake (X)、Snake (Y)")
P("另外 4 项（Slice/Divide Line、Slice/Divide Circle、Center Mark、Centerline）严格来说属于编辑/标注工具，"
  "不在「直线绘制」范畴，本次作为已知留白处理。")

H("二、核心设计原则", 1)
bullet("几何与 UI 解耦：新增 src/core/LineModes.h/.cpp，所有画法只接受点、参考图元 ID、光标位置，返回线段端点；"
       "视图层和命令行共用同一套计算，避免两份实现。")
bullet("activeTool 标识扩展：每个画法对应一个独立的工具 id（line_2pt / line_h / line_v / line_parallel / ...），"
       "与原有 line 工具互不干扰。")
bullet("参考图元复用 pickedIds_：已有的修剪/延伸两步拾取缓存被扩展为「参考图元收集器」，"
       "用于 Parallel / Tangent / Bisector 等需要先选参考线的模式。")
bullet("向后兼容：原 line 命令仍表示 2 Points；主工具条上的「直线」按钮默认进入 line，下拉菜单进入各子模式。")

H("三、模式分类与交互流程", 1)
P("按所需输入把 18 种模式分为四类：", bold=True)

tbl = doc.add_table(rows=1, cols=4)
tbl.style = 'Light Grid Accent 1'
hdr = tbl.rows[0].cells
hdr[0].text = "分类"
hdr[1].text = "代表模式"
hdr[2].text = "输入步骤"
hdr[3].text = "计算要点"
rows = [
    ("两点直接", "2 Points / Horizontal / Vertical / Angle", "1) 起点 2) 终点/长度参考点", "两点直接投影或射线角度约束"),
    ("单参考线", "Parallel / Orthogonal / Relative angle / From Point To Line / Orthogonal From Line", "1) 选参考线 2) 点...", "从参考线提取方向，做平行/垂直/定角"),
    ("双参考", "Bisector / Tangent (C,C) / Tangent Orthogonal", "1) 参考1 2) 参考2 3) 定位点", "两线求交、两圆公切线、圆+线垂切"),
    ("多点折线", "Snake / Snake(X) / Snake(Y)", "连续点击，右键/Enter 结束", "每段正交交替，X/Y 控制第一段方向"),
]
for a,b,c,d in rows:
    cells = tbl.add_row().cells
    cells[0].text = a; cells[1].text = b; cells[2].text = c; cells[3].text = d

doc.add_paragraph()

H("四、关键几何实现", 1)
P("4.1 参考线提取", bold=True)
P("对直线直接取 pts[0]-pts[1]；对矩形返回离光标最近的边所在无限直线。这让矩形边也能充当参考。")
code("static bool getRefLine(const Element& e, const Vec3& nearPoint, Vec3& a, Vec3& b)")

P("4.2 切线计算", bold=True)
P("Tangent (P,C)：利用经典公式 φ=atan2(P−C)，α=acos(r/|P−C|)，切点方向 φ±α。"
  "对圆弧额外用 angleInSweep 过滤，保证切点落在弧段上。")
P("Tangent (C,C)：将外公切问题转化为「切线段与连心线夹角 γ，满足 sinγ=(r1−r2)/D」。"
  "两条外公切线通过光标侧选中。")
P("Tangent Orthogonal：切线垂直于参考线 ⇒ 半径平行于参考线 ⇒ 切点为 C ± r·dir，再按光标选侧。")

P("4.3 Snake 折线", bold=True)
P("维护一个「下一段方向」状态。当 pts.size()==1 时由 Snake/SnakeX/SnakeY 决定首段方向；"
  "之后每完成一段就翻转方向。这样无需额外状态变量，纯由已有点序列推导。")
code("bool horizontal = (pts.size()==1) ? xFirst : !lastWasHorizontal;\n"
     "if (horizontal) return {last, Vec3(cur.x, last.y, 0)};\n"
     "else            return {last, Vec3(last.x, cur.y, 0)};")

P("4.4 角平分线", bold=True)
P("先求两无限直线交点 IP，再取两方向单位向量之和/差得到两条角平分线方向，"
  "按定位点 p 所在侧和光标距离选取最终端点。")

H("五、UI 层改动", 1)
P("5.1 工具条", bold=True)
P("将原先单独的「直线」按钮替换为 QToolButton::MenuButtonPopup：")
bullet("主按钮：line（2 Points），保留 L 快捷键")
bullet("下拉菜单：18 个子模式 QAction，全部加入同一 QActionGroup，实现单选高亮互斥")
P("5.2 菜单栏", bold=True)
P("在「草图绘制 > 直线绘制方式」子菜单中复用同一批 QAction，保证工具条、菜单、命令行三入口状态一致。")
P("5.3 视图交互", bold=True)
bullet("handlePress 优先判断 isLineSubTool(t)，避免普通 line 的两点逻辑干扰")
bullet("handleLineSubTool 分三步：收集参考图元 → 收集定位点 → 调用 line::compute 生成并 addLineSegment")
bullet("updatePreview 对 line_xxx 调用 line::compute，用当前光标作为 cur 实时反馈")
bullet("Backspace 现在可以回退参考图元；Snake 用右键或 Enter 结束")
P("5.4 快捷键重映射", bold=True)
P("isDrawTool 把 line_xxx 也识别为绘制命令，因此进入任何直线子模式时都会自动保存/恢复冲突快捷键，"
  "E/T/M/C/X 捕捉切换键在子模式下同样生效。")

H("六、命令行入口", 1)
P("在 CommandInterpreter 中统一处理 cmd.startsWith(\"line_\")，语法示例：")
code("line_2pt 0,0 100,50\n"
     "line_h 0,0 100,50\n"
     "line_angle 0,0 100,0 30\n"
     "line_parallel 1 0,20 80,20      # refId point1 point2\n"
     "line_from_point_to_line 1 30,40 # refId externalPoint\n"
     "line_tangent_pc 1 60,0          # circleId externalPoint\n"
     "line_tangent_cc 1 2             # circleId1 circleId2\n"
     "line_snake 0,0 50,0 50,30 80,30")
P("命令补全表同步扩展，前缀 li 现在会列出 line 及所有 line_xxx。")

H("七、测试覆盖", 1)
P("新增 selftest §22，覆盖：")
bullet("水平线 / 竖直线：坐标不变性")
bullet("Angle 模式：30° 射线角度校验")
bullet("Parallel / Orthogonal：与参考线平行/垂直")
bullet("From Point To Line：垂足正确性")
bullet("Tangent (P,C)：切点落在圆上")
bullet("Tangent (C,C)：等径圆水平外公切线")
bullet("Snake(X)：第一段为水平线")
bullet("CLI：line_h / line_v 执行成功")
P("全量自测 167 项通过 / 0 失败，构建 0 警告。")

H("八、代码落点", 1)
bullet("src/core/LineModes.h / .cpp：模式枚举、compute 分发器、各原子几何构造")
bullet("src/ui/GraphicsView.h / .cpp：isLineSubTool / handleLineSubTool / addLineSegment / finalizeSnake / pickRefElement；"
       "updatePreview、keyPressEvent、handleRightButton 的 Snake 与回退支持")
bullet("src/ui/MainWindow.cpp：直线 QToolButton 下拉菜单 + 草图绘制菜单子项")
bullet("src/core/CommandInterpreter.cpp：line_xxx 命令解析与执行")
bullet("CMakeLists.txt：LineModes.cpp 加入 SketchCAD 与 SketchSelfTest 两个目标")
bullet("tests/selftest.cpp：§22 直线子模式回归")

H("九、已知简化与后续扩展", 1)
bullet("Slice/Divide Line、Slice/Divide Circle、Center Mark、Centerline 不属于直线绘制，暂未实现；"
       "若后续需要可作为「编辑」或「标注」工具单独扩展。")
bullet("Angle / Relative angle 当前使用固定默认角度 30°；后续可在属性面板或命令行提供动态角度输入框。")
bullet("Tangent (C,C) 在不等径圆时生成的是外公切线；内公切线可作为额外选项加入。")
bullet("Snake 目前生成多条独立 Line 元素；后续可改为单个 Polygon 元素以支持整体编辑。")

# 页脚时间
P("")
pf = P("文档生成时间：2026-08-05", size=9)
pf.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

doc.save(DOC)
print(f"Saved: {DOC}  ({os.path.getsize(DOC)} bytes)")
