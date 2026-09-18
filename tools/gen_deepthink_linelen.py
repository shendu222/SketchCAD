# -*- coding: utf-8 -*-
"""生成《直线标注显示真实长度（修复投影长度）深度思考》Word 文档。
对应一次关于「标注直线时显示坐标轴投影长度而非真实长度」bug 的修复分析：
根因定位、修复方案、数学证明、验证策略、边界与取舍，以及它与更大范围
「标注交互/数值可编辑/属性联动」需求的关系。
注意：中文文案字符串统一用单引号包裹，避免与中文引号混淆。
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "宋体"
CODE_FONT = "Consolas"
doc = Document()

def set_cjk(style_name, size=None, bold=None):
    st = doc.styles[style_name]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold

set_cjk("Normal", 11)
for hs, sz in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
    set_cjk(hs, sz)

def para(text='', style=None, bold=False, italic=False, size=None, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    for ln in text.split('\n'):
        r = p.add_run(ln + '\n')
        r.font.name = CODE_FONT
        r._element.rPr.rFonts.set(qn('w:eastAsia'), CODE_FONT)
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return p

def add_figure(path, caption, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.font.name = FONT
    cr._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    cr.italic = True
    cr.font.size = Pt(10)
    cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()

# ===================== 标题 =====================
para('直线标注缺陷修复：从「坐标轴投影长度」到「真实长度」', style='Title', align=WD_ALIGN_PARAGRAPH.CENTER)
para('深度思考：根因定位 · 修复方案 · 数学证明 · 验证策略 · 边界取舍',
     style='Subtitle', align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=RGBColor(0x66, 0x66, 0x66))
para('基于 sketch-cad-cpp 项目（C++17 / Qt6）的标注子系统（DimStyle / SketchModel / GraphicsView）',
     size=9.5, color=RGBColor(0x99, 0x99, 0x99), align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

# ===================== 一、问题现象 =====================
para('一、问题现象与需求澄清', style='Heading 1')
para('用户反馈：在草图模块里对一条直线执行标注时，系统显示的数值并不是这条直线的实际长度，'
     '而是「直线在坐标轴上的投影长度」。例如一条从 (0,0) 到 (3,4) 的直线，其真实长度为 5（欧氏距离），'
     '但程序标出了 4 —— 这恰好是它在 Y 轴上的投影 |Δy|（因为 |Δy|=4 > |Δx|=3，取了较长的那个投影）。')
para('需求因此被澄清为：标注直线时，标注对象应是直线的「长度」，即两端点之间的欧氏距离；'
     '坐标轴投影长度（Δx / Δy）只在用户显式要求「水平标注 / 竖直标注」时才适用。')

# ===================== 二、根因定位 =====================
para('二、根因定位（为什么显示成投影）', style='Heading 1')
para('项目里标注的「实测值」有两处计算，但都正确区分了类型——问题不在计算，而在「直线被错误地'
     '分配成了投影类标注类型」。关键代码链路如下：')
bullet('标注类型枚举 DimensionType 含 Distance（线性自动）、Aligned（对齐）、Horizontal（水平）、'
       'Vertical（竖直）等。其中：')
code('enum class DimensionType {\n'
     '    Distance = 0,   // 线性（自动判水平/竖直，取投影较长的方向）\n'
     '    Aligned,        // 对齐（尺寸线平行于两测量点连线，量真实距离）\n'
     '    Horizontal,     // 线性-水平（只量 Δx）\n'
     '    Vertical        // 线性-竖直（只量 Δy）\n'
     '};')
bullet('测量函数 geo::measureDimension 对 Distance 的实现就是「取投影较长的方向」：')
code('case DimensionType::Distance: {\n'
     '    const double dx = std::fabs(e.pts[1].x - e.pts[0].x);\n'
     '    const double dy = std::fabs(e.pts[1].y - e.pts[0].y);\n'
     '    return dx >= dy ? dx : dy;   // 自动线性：取投影较长的方向\n'
     '}')
bullet('而 Aligned 的实现是真实长度：segLength = hypot(Δx, Δy)。两者语义本就不同。')
para('真正的问题在「分配类型」这一步。用户最常用的「标注」按钮在工具表里映射为 "dim" / "dim_linear"，'
     '而 dimToolType 把这两个工具解析为 DimensionType::Distance（见 src/ui/GraphicsView.cpp）：')
code('if (tool == "dim" || tool == "dim_linear") return DimensionType::Distance;')
para('随后在 handleDimension 的「线性族」分支里，点选到直线后直接把工具类型透传成了标注类型：')
code('// 修改前\n'
     'Element el = newElement(ElementType::Dimension, tname);\n'
     'el.dtype = dt;            // dt 来自工具，默认 dim→Distance\n'
     'el.pts = {A, B, pos};     // A、B 是直线的真实端点')
para('结论：直线被建成了一个 Distance 标注，于是 measureDimension 走「取较长投影」分支，'
     '对斜线给出投影长度而非真实长度。这就是 bug 根因——本质是「标注类型选择」与「图元语义」错配，'
     '而非几何计算错误。')

# ===================== 三、修复方案 =====================
para('三、修复方案：把「直线源」归一为对齐标注', style='Heading 1')
para('修复核心思想：当标注的源图元是「直线」时，除非用户显式要求水平/竖直，否则一律将其标注类型'
     '归一为 Aligned（量真实长度）。这样既不改 Distance 类型本身的投影语义（手动两点距离标注仍可用），'
     '又让「点直线出标注」这一最自然的操作得到正确结果。')
para('为避免 UI 与测试各写一份判断，把归一逻辑抽成几何层纯函数 geo::lineDimType，UI 与自测共用，'
     '保证决策唯一、可测：', style='Heading 3')
code('// src/core/Geometry.h（内联，几何层，无 UI 依赖）\n'
     'inline DimensionType lineDimType(DimensionType requested) {\n'
     '    return (requested == DimensionType::Horizontal || requested == DimensionType::Vertical)\n'
     '               ? requested : DimensionType::Aligned;\n'
     '}')
para('UI 侧 handleDimension 的直线分支改为使用它，并把归一后的类型用于标签与日志：', style='Heading 3')
code('// src/ui/GraphicsView.cpp —— 直线分支\n'
     'const DimensionType useDt = geo::lineDimType(dt);  // 直线→强制对齐（量真实长度）\n'
     'const QString useName = QString::fromUtf8(dimensionTypeName(useDt));\n'
     'Element el = newElement(ElementType::Dimension, useName);\n'
     'el.dtype = useDt;\n'
     'el.refA = e->id;     // 关联源图元 → 标注数值可驱动几何\n'
     'el.pts = {A, B, pos};')
para('另一条入口 dim_smart（智能标注）本来就用 smartDimension() 把直线建成 Aligned，已正确，'
     '本次无需改动，修复后两条入口语义一致。')

# ===================== 四、数学证明 =====================
para('四、为什么 Aligned 一定等于真实长度', style='Heading 1')
para('Aligned 的尺寸线方向定义为「沿被测两点连线」的单位向量：')
code('case DimensionType::Aligned: dir = normalize(sub(B, A)); break;')
para('实测值在 buildLinear 里是投影到该方向上的标量长度：')
code('g.measured = std::fabs(dot2(sub(B, A), dir));')
para('设 d = B − A，dir = d / |d|，则 dot(d, dir) = d · (d/|d|) = |d| = hypot(Δx, Δy) = 真实长度。'
     '方向完全相同，投影即为全长，因此 Aligned 的 measured 恒等于欧氏距离。'
     '以 (0,0)→(3,4) 为例：dir=(0.6,0.8)，dot((3,4),(0.6,0.8))=1.8+3.2=5.0，正确。')
para('对比 Distance：dir 取 (0,1)（因为 |Δy|>|Δx|），dot((3,4),(0,1))=4.0，即投影长度——正是被修复的旧行为。')

# ===================== 五、验证策略 =====================
para('五、验证策略与自测回归', style='Heading 1')
para('修复后用无界面自测（tests/selftest.cpp §15 新增 6 项）锁定语义，避免回归：')
bullet('geo::measureDimension(Distance, (0,0)-(3,4)) == 4.0：确认 Distance 仍是投影行为（旧错误的来源，作为对照保留）。')
bullet('geo::lineDimType(Distance) == Aligned、lineDimType(Aligned) == Aligned：归一决策正确。')
bullet('geo::lineDimType(Horizontal) == Horizontal、lineDimType(Vertical) == Vertical：显式水平/竖直不被覆盖。')
bullet('用 lineDimType(Distance) 归一后 build 直线标注，measured == 5.00 且 label == "5.00"：修复后点选直线得到真实长度。')
para('全量自测结果：156 通过 / 0 失败，构建 0 警告。')
para('运行期人工验证：启动 build/SketchCAD.exe，画一条斜线（如从 (0,0) 拖到 (3,4)），'
     '按「标注」按钮后直接点选该直线，标注数字应显示 5.00；按 L 切换显式「水平标注」再点同一直线，'
     '则显示 3.00（Δx），符合预期。')

# ===================== 六、边界与取舍 =====================
para('六、边界情况与取舍', style='Heading 1')
bullet('显式水平/竖直标注：保留投影语义。这是合理的——用户主动要 Δx/Δy 时，标注投影长度才是意图。')
bullet('手动两点「距离标注」：点两自由点（非点选直线）的 dim 仍走 Distance 投影语义，符合「线性尺寸自动选水平/竖直」的国标习惯；本修复只针对「以直线为源」这一情形。')
bullet('dim_smart 智能标注：直线→Aligned 本已正确，修复后两条入口语义一致，无行为回退。')
bullet('关联驱动：修复同时保留 el.refA = e->id，使标注与源直线关联，为后续「改标注数值→驱动直线长度」留好钩子（见下一节）。')
bullet('容差/偏置：尺寸线偏置方向由点击点相对直线的垂距决定，与采用哪种线性类型无关，不受本次改动影响。')

# ===================== 七、与更大范围需求的关系 =====================
para('七、与「标注交互 / 数值可编辑 / 属性联动」需求的关系', style='Heading 1')
para('本次修复直接回应了「标注直线应显示真实长度」这一具体诉求，属于「标注正确性」层面的修正。'
     '它也是后续更大范围需求的前置基础：')
bullet('【标注功能交互优化】「点击标注按钮 → 直接选图元 → 标注」的无冗余流程，已在代码中以 '
       'handleDimension 直接命中图元实现（代码中已标「需求 1 交互优化」），与本次类型修复正交、互不影响。')
bullet('【标注尺寸数值可编辑 + 联动】要让「改标注数值 → 直线按新长度自动缩放、相连图元联动」，'
       '需要把标注实测值作为驱动源，反向求解 refA 图元的端点（保持方向、改长度），并沿「连接关系」传播到相邻图元。'
       '这要求引入轻量参数/约束求解层；当前 el.refA 已为「标注→图元」建立关联，是第一步。')
bullet('【图元属性联动更新】修改某图元属性时让相连图元同步，本质是同一套「几何依赖图 + 增量求解」机制。')
para('结论：本次交付聚焦且自洽——先纠正确实影响正确性的「投影长度」缺陷（已修复、已测、已文档化）；'
     '更大范围的参数化/联动需求建议作为后续独立任务推进，可复用已建立的 refA 关联与几何层接口。')

# ===================== 输出 =====================
out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'docs')
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, '直线标注真实长度修复-深度思考.docx')
doc.save(out_path)
print('saved:', out_path)
